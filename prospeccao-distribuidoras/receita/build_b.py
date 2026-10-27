import json, re, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/home/user/comercial/prospeccao-distribuidoras/lista-B-onda1-30-distribuidoras-novas.xlsx"
S = {x["base"]: x for x in json.load(open("out/shortlist.json"))}
CNAE = {c: d for c, d in csv.reader(open("rf_ref/cnae_utf8.csv", encoding="utf8"), delimiter=";")}
PICK = ["11989122","11247450","05265262","43252675","06232652","18871890","07860891","40755087","26185195","15253243",
        "08562964","15104520","45985626","04588373","08539133","04575811","25449195","44298564","34041526","01181512",
        "04574350","42089430","38471161","36703258","22248811","12983291","10809205","02612347","11670333","27665636"]
CLIENTES = {"01611823":"Modenuti (ativo, R$ 9,1 mi)","05951362":"Centerlar (ativo, R$ 4,2 mi)","04150555":"Millenium (ativo, R$ 4,0 mi)",
            "11209781":"ABS / Palmares (grupo ativo)","26361858":"Parolar (ativo)","04043178":"Cordeirópolis UD / Campineira (ativo)",
            "23030412":"Guará / Campineira (ativo)","67742288":"Soroca 2 (em risco)","01387543":"SPF (perdido, está na lista A)",
            "18377927":"BBB Utilidades (perdido)"}
CAD_SEM_COMPRA = {"07063181","02184384","11022033","65484453","04575811","10809205"}
FORA = {"52321224":"Parece loja de shopping (varejo)","25301503":"Foco em cama, mesa e banho/decoração","52459079":"Parece fábrica (Plásticos MB): possível concorrente","11903765":"Ligada à Isoforma (fabricante de descartáveis)",
        "03640671":"Plasnew: fabricante/marca concorrente","11071732":"Bettanin: fabricante concorrente","04381287":"Ordene: fabricante concorrente",
        "35548485":"Quality House: marca própria importada (concorrente)","43670056":"Cresci e Perdi: franquia de usados, fora do perfil"}
OBS = {
 "11989122":"Nome e CNAE de distribuidora de UD; 16 anos. Cara de Modenuti/Millenium.",
 "11247450":"Opera como Marcamix (importadora). Confirmar se revende outras marcas.",
 "43252675":"Atacadão de presentes e utilidades, 47 anos.",
 "06232652":"Distribuidora de UD com 23 anos em Itapetininga.",
 "04575811":"Já tem CADASTRO na Nitron (dez/22), mas nunca comprou nos últimos 24 meses: retomar.",
 "10809205":"Central de distribuição de rede (Bom Lugar). Já tem CADASTRO na Nitron, sem compra em 24 meses.",
 "01181512":"Telefone da Receita inválido: buscar no Google/site.",
 "04574350":"Sem telefone na Receita: buscar no Google.",
 "08539133":"Pedreira é polo de fábricas de plástico: confirmar se revende ou fabrica.",
 "02612347":"Região Central = área da Modenuti. Alinhar com o Sebastião antes.",
 "04574350":"Sem telefone na Receita. Região Central = área da Modenuti: alinhar com o Sebastião.",
 "27665636":"Usa o nome Favorito Atacado. Região Central = área da Modenuti: alinhar com o Sebastião.",
 "11670333":"Distribuidora Maestral (domínio próprio); 17 anos.",
 "11247450":"Opera como Marcamix/Houseware (importadora com marca própria): confirmar se revende outras marcas ou se é concorrente.",
}
def fmt_cnpj(c): return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"
def contador(e): return bool(re.search(r"contab|contador|escritorio|adv@|\.adv", e))
def tel_ok(t): return bool(t) and not re.search(r"1111111|0000000", t)
PORTE = {"03":"EPP","05":"Demais (médio/grande)","01":"ME"}

FONT="Arial"; H_FILL=PatternFill("solid", fgColor="1F3864"); H_FONT=Font(name=FONT,bold=True,color="FFFFFF",size=10)
B_FONT=Font(name=FONT,size=10); IN_FILL=PatternFill("solid", fgColor="FFF2CC"); thin=Side(style="thin",color="D9D9D9")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
def header(ws, cols, widths, freeze="C2"):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell=ws.cell(row=1,column=i,value=c); cell.fill,cell.font,cell.border=H_FILL,H_FONT,BORDER
        cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=32; ws.freeze_panes=freeze

def ordem(x):  # mais aderente primeiro: pontuacao, porte, idade
    return (-x["score"], x["porte"]!="05", -x["idade"])

COLS=["#","Razão social","Nome fantasia","CNPJ","Cidade","Região de entrega","Endereço","Bairro","CEP","Telefone 1","Telefone 2",
      "E-mail (Receita)","Atividade principal (CNAE)","Porte","Capital social (R$)","Anos de empresa","Pontuação de aderência","Observação",
      "Status","Data do contato","Responsável","Decisor (nome/cargo)","Telefone/WhatsApp correto","Compra de quais marcas hoje?","Próximo passo","Valor do 1º pedido (R$)"]
WID=[4,40,24,19,18,18,36,18,10,15,15,30,40,14,13,9,10,46,16,12,14,22,18,26,28,14]
STATUS=["A contatar","Sem contato","Número errado","Em conversa","Cadastro enviado","Pedido fechado","Sem interesse","Não é distribuidor"]

wb=Workbook(); leia=wb.active; leia.title="Leia-me"
ws=wb.create_sheet("30 distribuidoras"); header(ws,COLS,WID)
sel=sorted((S[b] for b in PICK), key=ordem)
for r,x in enumerate(sel,2):
    obs=[OBS.get(x["base"],"")]
    if x["base"] in CAD_SEM_COMPRA and "CADASTRO" not in obs[0]: obs.append("Já tem cadastro na Nitron, sem compra em 24 meses.")
    if contador(x["email"]): obs.append("E-mail parece ser do contador.")
    if not tel_ok(x["tel1"]) and "Telefone" not in obs[0] and "telefone" not in obs[0]: obs.append("Telefone da Receita ausente/inválido.")
    vals=[r-1,x["razao"],x["fantasia"] or "",fmt_cnpj(x["cnpj"]),x["cidade"].title(),x["regiao"],re.sub(r"\s+"," ",x["end"]).title(),
          x["bairro"].strip().title(),f"{x['cep'][:5]}-{x['cep'][5:]}",x["tel1"] if tel_ok(x["tel1"]) else "",x["tel2"],x["email"],
          f"{x['cnae']} - {CNAE.get(x['cnae'],'')}",PORTE.get(x["porte"],x["porte"]),x["cap"],round(x["idade"]),x["score"]," ".join(o for o in obs if o),
          "A contatar",None,None,None,None,None,None,None]
    for c,v in enumerate(vals,1):
        cell=ws.cell(row=r,column=c,value=v); cell.font,cell.border=B_FONT,BORDER
        cell.alignment=Alignment(vertical="top",wrap_text=c in (2,3,7,13,18,24,25))
        if c>=19: cell.fill=IN_FILL
    ws.cell(row=r,column=15).number_format='#,##0;(#,##0);"-"'
    ws.cell(row=r,column=20).number_format="DD/MM/YYYY"
    ws.cell(row=r,column=26).number_format='#,##0;(#,##0);"-"'
last=len(sel)+1
dv=DataValidation(type="list",formula1='"'+",".join(STATUS)+'"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"S2:S{last}")
ws.auto_filter.ref=f"A1:{get_column_letter(len(COLS))}{last}"

# reserva
NEG=r"BRINQUED|TOYS|KIDS|SKATE|FESTA|PAPEL DE PAREDE|RELIGIOS|CLIPPER|CHEMICAL|LEGO|MUSICA|DRUMS|BALLOON|FOSFORO|PERFUMARIA|COMUNICACAO|SUZANO|MAKRO|TENDA|FASTENAL|DESTRO|SCAPOL|ALIMENT|PET"
res=[x for b,x in S.items() if b not in PICK and b not in CLIENTES and b not in FORA and not re.search(NEG,x["razao"]+" "+x["fantasia"])]
res.sort(key=ordem)
ws=wb.create_sheet("Reserva (não revisada)")
rcols=["Razão social","Nome fantasia","CNPJ","Cidade","Região de entrega","Telefone 1","E-mail (Receita)","Atividade principal (CNAE)","Porte","Anos de empresa","Pontuação de aderência"]
header(ws,rcols,[42,24,19,18,18,15,30,40,14,9,10],freeze="B2")
for r,x in enumerate(res,2):
    vals=[x["razao"],x["fantasia"] or "",fmt_cnpj(x["cnpj"]),x["cidade"].title(),x["regiao"],x["tel1"] if tel_ok(x["tel1"]) else "",x["email"],
          f"{x['cnae']} - {CNAE.get(x['cnae'],'')}",PORTE.get(x["porte"],x["porte"]),round(x["idade"]),x["score"]]
    for c,v in enumerate(vals,1):
        cell=ws.cell(row=r,column=c,value=v); cell.font,cell.border=B_FONT,BORDER; cell.alignment=Alignment(vertical="top",wrap_text=c in (1,8))
ws.auto_filter.ref=f"A1:K{len(res)+1}"

ws=wb.create_sheet("Fora da lista")
header(ws,["CNPJ (raiz)","Empresa","Por que saiu"],[14,50,60],freeze="A2")
r=2
for b,m in list(CLIENTES.items())+list(FORA.items()):
    nome=S[b]["razao"] if b in S else ""
    for c,v in enumerate([b,nome,("Já é cliente Nitron: " if b in CLIENTES else "")+m],1):
        cell=ws.cell(row=r,column=c,value=v); cell.font,cell.border=B_FONT,BORDER; cell.alignment=Alignment(vertical="top",wrap_text=True)
    r+=1

leia.column_dimensions["A"].width=30; leia.column_dimensions["B"].width=100
leia["A1"]="Lista B · Onda 1: 30 distribuidoras que a Nitron ainda não atende"; leia["A1"].font=Font(name=FONT,bold=True,size=14,color="1F3864")
leia["A2"]="Fonte: dados abertos de CNPJ da Receita Federal (base out/2026), cruzados com o cadastro do Sankhya em 27/10/2026."; leia["A2"].font=Font(name=FONT,italic=True,size=9,color="595959")
L=[
 ("Resumo",None),
 ("Distribuidoras selecionadas","=COUNTA('30 distribuidoras'!B2:B31)"),
 ("  com telefone na Receita","=COUNTIF('30 distribuidoras'!J2:J31,\"?*\")"),
 ("  já contatadas",f"=COUNTA('30 distribuidoras'!T2:T31)"),
 ("  pedidos fechados","=COUNTIF('30 distribuidoras'!S2:S31,\"Pedido fechado\")"),
 ("  valor dos 1ºs pedidos (R$)","=SUM('30 distribuidoras'!Z2:Z31)"),
 ("Reserva para a próxima rodada",f"=COUNTA('Reserva (não revisada)'!A2:A{len(res)+1})"),
 (None,None),
 ("Como a lista foi feita",None),
 ("Região","Interior de SP sem adicional de frete CIF: Campinas, Extremo de Campinas, Sorocaba, Extremo de Sorocaba, Região Central e Vale do Paraíba (198 cidades)."),
 ("Funil","4.941 empresas ativas de atacado nessas cidades → 1.722 com porte EPP ou maior, 3+ anos e sem ser MEI → 170 com perfil de distribuidor de utilidades → 16 já estão no cadastro Nitron → 30 escolhidas a dedo."),
 ("Atividades (CNAE)","Atacado de artigos de uso doméstico (4649-4/99), mercadorias em geral (4693-1/00), embalagens (4686-9/02), higiene e limpeza (4649-4/08) e ferragens (4672-9/00)."),
 ("Pontuação de aderência","Soma pontos por atividade de utilidades domésticas, nome com utilidades/distribuidora/atacado, porte e tempo de empresa. Teste: as 3 empresas de maior pontuação da região são Modenuti, Centerlar e Millenium, as 3 maiores distribuidoras da Nitron."),
 (None,None),
 ("Antes de ligar",None),
 ("Telefone","É o declarado na Receita: pode estar velho ou ser do contador. Se não atender, buscar no Google Maps/site. Anotar o número certo na coluna amarela."),
 ("Território","Modenuti (São Carlos), Millenium (Campinas/Salto), Centerlar (Jundiaí) e Parolar (Saltinho) ficam dentro desta região. Alinhar com Sebastião e João Carlos se a prospecção pode concorrer com eles."),
 ("Colunas amarelas","Para o time preencher. Status tem lista suspensa. Exemplo: Status = Em conversa · Decisor = João, comprador · Compra de = Sanremo, Plasútil · Próximo passo = enviar tabela de distribuidor até 05/11."),
 (None,None),
 ("Roteiro da ligação fria",None),
 ("1. Abertura","\"Oi, aqui é [nome] da Nitron, fabricante de potes, organizadores e utilidades. Falo com quem cuida das compras?\""),
 ("2. Qualificação","Vocês distribuem utilidades domésticas para lojas/mercados da região? Quais marcas trabalham hoje? Qual o volume de compra por mês? Atendem quais cidades?"),
 ("3. Gancho","Nossas linhas que mais giram nas distribuidoras da região são Organização e Potes. Frete CIF sem adicional para a sua cidade, pedido mínimo de R$ 2.500 (R$ 3.570 nos 'Extremos')."),
 ("4. Fechamento","Sair com o envio da tabela + cadastro, e uma data para retorno. Nunca terminar sem próximo passo."),
]
row=4
for k,v in L:
    if k is not None:
        a=leia.cell(row=row,column=1,value=k); a.font=Font(name=FONT,size=11,bold=True,color="1F3864") if v is None else Font(name=FONT,size=10)
    if v is not None:
        b=leia.cell(row=row,column=2,value=v); b.font=B_FONT; b.alignment=Alignment(wrap_text=True,vertical="top")
        if v.startswith("=SUM"): b.number_format='#,##0;(#,##0);"-"'
    row+=1
wb.save(OUT); print(OUT, len(sel), len(res))
